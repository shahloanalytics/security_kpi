"""Runtime Excel/PDF report generation from immutable snapshots."""
import io
from pathlib import Path
from xml.sax.saxutils import escape
from metrics import STATES
from imports import safe_value

COLUMNS=[('code','KPI'),('name','Ko‘rsatkich'),('plan','Reja'),('actual','Amalda'),('rate','Bajarilish %'),('target','Maqsad %'),('score','KPI balli'),('weight','Asl vazn %'),('effective_weight','Amaldagi vazn %'),('contribution','Vaznli hissa'),('state','Holat')]
def display(v):return '—' if v is None else f'{v:.2f}' if isinstance(v,float) else str(v)
def excel_report(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.chart import BarChart,Reference
    from openpyxl.utils import get_column_letter
    s=report['payload'];calc=s['calculation'];wb=Workbook();ws=wb.active;ws.title='Hisobot'
    meta=[['Security KPI — Avtozavod',''],['Tashkilot',s['org']],['Davr',s['month']],['Baholash',s['entity']],['Metodika',s['template']['id']],['Metodika holati',s['template']['label']],['Hisobot holati',report['state']],['Natija',calc['total'] if calc['final'] else 'Dastlabki hisob'],['Kuzatilgan vaznli ball',calc['observed_points']],['To‘liqlik %',calc['coverage']],['Muallif',report['creator']],['Tekshiruvchi',report.get('reviewer') or 'Belgilanmagan'],['Tasdiqlovchi',report.get('approver') or 'Belgilanmagan'],['Hisobot ID',report['id']]]
    for row in meta:ws.append([safe_value(x) for x in row])
    detail=wb.create_sheet('KPI_hisobi');detail.append([label for _,label in COLUMNS])
    for r in calc['rows']:detail.append([safe_value(STATES.get(r[k],r[k]) if k=='state' else r.get(k)) for k,_ in COLUMNS])
    # Frozen report values intentionally do not recalculate after approval.
    chart=BarChart();chart.title='KPI ballari (0–100)';chart.add_data(Reference(detail,min_col=7,min_row=1,max_row=len(calc['rows'])+1),titles_from_data=True);chart.set_categories(Reference(detail,min_col=1,min_row=2,max_row=len(calc['rows'])+1));chart.y_axis.scaling.max=100;ws.add_chart(chart,'D2')
    actions=wb.create_sheet('Chora_tadbirlari');actions.append(['Vazifa','Mas’ul','Muddat','Holat'])
    for a in s['actions']:actions.append([safe_value(a[k]) for k in ('task','owner','due','state')])
    issues=wb.create_sheet('Tahlil');issues.append(['Kuzatuv va tavsiya'])
    for line in s['analysis']:issues.append([safe_value(line)])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='153E4D');cell.font=Font(color='FFFFFF',bold=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,float):cell.number_format='0.00'
        for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=22 if col!=2 else 45
        for row in range(2,sheet.max_row+1):sheet.row_dimensions[row].height=42
    issues.column_dimensions['A'].width=110
    b=io.BytesIO();wb.save(b);return b.getvalue()

def pdf_report(report):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.graphics.shapes import Drawing,Rect,String
    font=Path(__file__).parent/'assets'/'DejaVuSans.ttf'
    if 'UZ' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('UZ',str(font)))
    styles=getSampleStyleSheet()
    for style in styles.byName.values():style.fontName='UZ'
    styles.add(ParagraphStyle(name='SmallUz',fontName='UZ',fontSize=8,leading=11,spaceAfter=5))
    b=io.BytesIO();doc=SimpleDocTemplate(b,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=40)
    s=report['payload'];calc=s['calculation'];story=[]
    def p(text,style='SmallUz'):return Paragraph(escape(str(text)),styles[style])
    story += [p('Security KPI — Avtozavod','Title'),p('Mustaqil namuna platforma. BYD rasmiy mahsuloti emas.'),Spacer(1,10)]
    for k,v in [('Tashkilot',s['org']),('Davr / baholash',s['month']+' / '+s['entity']),('Metodika',s['template']['id']+' · '+s['template']['label']),('Hisobot',report['id']+' · '+report['state']),('To‘liqlik',f"{calc['coverage']:.1f}%"),('Natija',f"{calc['total']:.2f} ball" if calc['final'] else 'Dastlabki hisob — yakuniy reyting yo‘q')]:story.append(p(k+': '+v))
    story += [Spacer(1,10),p('KPI hisoblash jadvali','Heading2')]
    cols=[('code','KPI'),('plan','Reja'),('actual','Amal'),('rate','Bajar. %'),('target','Maqs. %'),('score','Ball'),('effective_weight','Vazn %'),('contribution','Hissa'),('state','Holat')]
    data=[[p(x[1]) for x in cols]]+[[p(STATES.get(r.get(k),r.get(k)) if k=='state' else display(r.get(k))) for k,_ in cols] for r in calc['rows']]
    t=Table(data,colWidths=[60,47,47,48,48,44,46,44,139],repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dfecef')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#ced8dc')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(t)
    story.append(p('Formula: bajarilish = amalda / reja × 100; ball = min(100, bajarilish / maqsad × 100); hissa = ball × amaldagi vazn / 100.'))
    story.append(p(f"Qo‘llanilmaydigan KPI vazni: {calc['na_weight']:.2f}%. Qayta taqsimlandi: {'ha' if calc['redistributed'] else 'yo‘q'}. Asl va amaldagi vaznlar Excel hisobotida alohida berilgan."))
    story += [PageBreak(),p('Natijalar tahlili','Heading1')]
    d=Drawing(510,len(calc['rows'])*25+20)
    for i,r in enumerate(calc['rows']):
        y=len(calc['rows'])*25-i*25
        d.add(String(0,y,r['code'],fontName='UZ',fontSize=9));d.add(Rect(80,y-3,350,13,fillColor=colors.HexColor('#e9eef0'),strokeColor=None))
        if r['score'] is not None:d.add(Rect(80,y-3,3.5*r['score'],13,fillColor=colors.HexColor('#257a70'),strokeColor=None))
        d.add(String(440,y,display(r['score']),fontName='UZ',fontSize=9))
    story.append(d)
    for line in s['analysis']:story.append(p(line))
    story += [Spacer(1,10),p('Vazifa — mas’ul — muddat','Heading2')]
    if s['actions']:
        ad=[[p(x) for x in ('Vazifa','Mas’ul','Muddat','Holat')]]+[[p(a[k]) for k in ('task','owner','due','state')] for a in s['actions']]
        at=Table(ad,colWidths=[270,85,85,83],repeatRows=1);at.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP')]));story.append(at)
    else:story.append(p('Chora-tadbir kiritilmagan. Vazifalar va muddatlarni mas’ul rahbar belgilaydi.'))
    story += [Spacer(1,15),p('Muallif: '+report['creator']),p('Tekshiruvchi: '+(report.get('reviewer') or '________________')),p('Tasdiqlovchi: '+(report.get('approver') or '________________'))]
    def footer(c,doc):c.setFont('UZ',8);c.drawString(36,22,'Security KPI | '+report['id']);c.drawRightString(559,22,str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return b.getvalue()

def analysis(calc):
    lines=[]
    for r in calc['rows']:
        if r['state'] not in ('ok','na'):lines.append(f"{r['code']}: {STATES[r['state']]}. Manba va dalillarni tekshiring.")
        elif r['score'] is not None and r['score']<100:
            gap=max(0,r['plan']*r['target']/100-r['actual'])
            lines.append(f"{r['code']} — {r['name']}: maqsadgacha {gap:.2f} {r['unit']} yetishmagan; yo‘qotilgan vaznli ball {r['effective_weight']-r['contribution']:.2f}. Sabablarni tekshirib, mas’ul va muddatli chora belgilang.")
        if r['plan'] is not None and 0<r['plan']<10:lines.append(f"{r['code']}: kuzatuv hajmi kichik ({r['plan']:g}); smenalar taqqoslashida ehtiyotkor talqin qiling.")
    if not lines:lines=['Hisoblanadigan mezonlarda maqsad bajarilgan. Muntazam monitoringni davom ettiring.']
    if not calc['final']:lines.insert(0,'Dastlabki hisob: ma’lumot/metodika tasdiqlanmaguncha yakuniy reyting berilmaydi.')
    return lines
